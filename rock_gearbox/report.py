"""选型计算书 Word 文档生成。"""

from __future__ import annotations
from datetime import datetime
from io import BytesIO

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

from .models import SelectionResult


def _set_default_font(doc: Document):
    style = doc.styles["Normal"]
    style.font.name = "Microsoft YaHei"
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    style.font.size = Pt(10.5)


def _add_heading(doc: Document, text: str, level: int = 1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.name = "Microsoft YaHei"
        run.element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)


def _add_kv_table(doc: Document, rows: list[tuple[str, str]]):
    table = doc.add_table(rows=len(rows), cols=2)
    table.style = "Light Grid Accent 1"
    for i, (k, v) in enumerate(rows):
        table.rows[i].cells[0].text = str(k)
        table.rows[i].cells[1].text = str(v)


def generate_report(
    res: SelectionResult,
    *,
    project_name: str = "",
    customer: str = "",
    engineer: str = "",
    image_png_bytes: bytes | None = None,
) -> bytes:
    """生成 Word 选型计算书,返回 .docx 字节流。"""
    doc = Document()
    _set_default_font(doc)

    # 标题
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Rock 齿轮箱选型计算书")
    run.font.size = Pt(22); run.font.bold = True
    run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run(f"型号:{res.type_code}").font.size = Pt(14)

    # 项目信息
    _add_heading(doc, "一、项目信息")
    _add_kv_table(doc, [
        ("项目名称", project_name or "(未填写)"),
        ("客户",     customer or "(未填写)"),
        ("工程师",   engineer or "(未填写)"),
        ("日期",     datetime.now().strftime("%Y-%m-%d")),
    ])

    # 输入工况
    _add_heading(doc, "二、输入工况")
    inp = res.inp
    _add_kv_table(doc, [
        ("电机功率 P₁", f"{inp.p1_kw} kW"),
        ("工作机功率 P₂", f"{inp.p2_kw} kW"),
        ("输入转速 n₁", f"{inp.n1_rpm} r/min"),
        ("输出转速 n₂", f"{inp.n2_rpm} r/min"),
        ("最大启动扭矩 T_A", f"{inp.t_a_nm or '由 P₂/n₂ 计算'} N·m"),
        ("工作机类型 / 应用",  inp.application_key),
        ("每小时工作周期 ED",  f"{inp.ed_pct} %"),
        ("每日运行小时数",     f"{inp.hours_per_day} h"),
        ("每小时启停次数",     f"{inp.starts_per_hour}"),
        ("负载方向",          {"unidirectional": "单向", "alternating": "交变"}[inp.load_direction]),
        ("原动机类型",        inp.prime_mover),
        ("传动类型(系列)",   {"RKH": "RKH 平行轴", "RKB": "RKB 直交轴", "auto": "不限"}[inp.family]),
        ("输出轴形式",        {"S": "实心轴", "H": "空心轴", "D": "空心轴+收缩盘"}[inp.output_shaft]),
        ("安装方位",          {"H": "卧式", "M": "立式"}[inp.mounting]),
        ("润滑方式",          {"oil_bath": "浸油润滑", "forced": "强制润滑"}[inp.lubrication]),
        ("环境温度",          f"{inp.ambient_temp_c} ℃"),
        ("散热方式偏好",       inp.cooling_pref),
    ])

    # 服务系数
    _add_heading(doc, "三、服务系数取值")
    _add_kv_table(doc, [
        ("f₁ — 工况系数",        f"{res.f1:.2f}  (Rock.pdf 表 1)"),
        ("f₂ — 原动机系数",      f"{res.f2:.2f}  (表 2)"),
        ("f₃ — 峰值扭矩系数",    f"{res.f3:.2f}  (表 3,基于 {inp.starts_per_hour} 启停/h, {inp.load_direction})"),
        ("f₄ — 温度系数(自然/风冷)", f"{res.f4:.2f}  (表 4,{inp.ambient_temp_c}°C, ED={inp.ed_pct}%)"),
        ("f₅ — 温度系数(水冷)",  f"{res.f5:.2f}  (表 5)"),
        ("f₈ — 供油系数",       f"{res.f8:.2f}  (表 6)"),
    ])

    # 计算
    _add_heading(doc, "四、选型计算")

    p = doc.add_paragraph()
    p.add_run("4.1 传动比\n").bold = True
    p.add_run(f"i_s = n₁ / n₂ = {inp.n1_rpm} / {inp.n2_rpm} = {res.i_required:.3f}\n")
    p.add_run(f"取最接近的标准 i_N = {res.chosen_ratio_nominal}\n")
    p.add_run(f"实际传动比 i = {res.chosen_ratio_actual:.3f}  "
              f"(偏差 {res.ratio_deviation_pct:.2f}%)\n")

    p = doc.add_paragraph()
    p.add_run("4.2 输出扭矩\n").bold = True
    p.add_run(f"T₂ = 9550 · P₂ / n₂ = 9550 × {inp.p2_kw} / {inp.n2_rpm} = {res.t_2_nm:.1f} N·m\n")

    p = doc.add_paragraph()
    p.add_run("4.3 所需额定功率\n").bold = True
    p.add_run(f"P_N ≥ P₂ · f₁ · f₂ = {inp.p2_kw} × {res.f1} × {res.f2} = {res.required_p_n_kw:.1f} kW\n")

    p = doc.add_paragraph()
    p.add_run("4.4 寿命/峰值扭矩校核\n").bold = True
    p.add_run(f"P_N ≥ T_A · n₁ · f₃ / 9550 = {res.t_a_nm:.0f} × {inp.n1_rpm} × {res.f3} / 9550 "
              f"= {res.required_p_n_life_kw:.1f} kW\n")

    # 推荐型号
    _add_heading(doc, "五、推荐型号铭牌")
    _add_kv_table(doc, [
        ("型号编码",         res.type_code),
        ("系列",            res.best.series_code),
        ("机座号",          str(res.best.size)),
        ("公称速比 i_N",     f"{res.best.ratio_nominal}"),
        ("实际速比 i",       f"{res.best.ratio_actual:.3f}"),
        ("额定输出扭矩 T₂N", f"{res.best.t2n_nm:.0f} N·m  ({res.best.t2n_nm/1000:.1f} kN·m)"),
        ("扭矩利用率",       f"{res.torque_utilization*100:.1f} %"),
    ])

    # 热功率校核
    _add_heading(doc, "六、热功率校核")
    if res.cooling_options:
        table = doc.add_table(rows=1 + len(res.cooling_options), cols=5)
        table.style = "Light Grid Accent 1"
        hdr = table.rows[0].cells
        for i, h in enumerate(["散热方式", "铭牌 P_G (kW)", "修正系数", "修正 P_G (kW)", "结论"]):
            hdr[i].text = h
        for i, c in enumerate(res.cooling_options):
            row = table.rows[i + 1].cells
            row[0].text = c.name_cn
            row[1].text = f"{c.p_g_rated_kw:.1f}"
            row[2].text = f"{c.f_temp_label}×f₈ = {c.f_temp:.2f}×{c.f8:.2f}"
            row[3].text = f"{c.p_g_corrected_kw:.1f}"
            row[4].text = "✓ 通过" if c.passes else "✗ 不足"
        p = doc.add_paragraph()
        p.add_run(f"\n选定散热方式:{res.chosen_cooling.name_cn}").bold = True
    else:
        doc.add_paragraph("(热功率数据未录入,本期跳过该校核)")

    # 校核结论
    _add_heading(doc, "七、校核结论")
    checks = [
        ("扭矩校核 T₂N ≥ T₂·f₁·f₂",  res.torque_check_pass),
        ("功率上限校核 3.33·P₂ ≥ P_N", res.power_3_33_check_pass),
        ("寿命校核",                  res.life_check_pass),
        ("热功率校核",                res.chosen_cooling.passes if res.cooling_options else True),
    ]
    table = doc.add_table(rows=len(checks), cols=2)
    table.style = "Light Grid Accent 1"
    for i, (label, ok) in enumerate(checks):
        table.rows[i].cells[0].text = label
        table.rows[i].cells[1].text = "✓ 通过" if ok else "✗ 不通过"

    # 备注
    if res.notes:
        _add_heading(doc, "八、备注")
        for n in res.notes:
            doc.add_paragraph(f"• {n}")

    # 外形图
    if image_png_bytes:
        _add_heading(doc, "附:外形/安装尺寸图")
        doc.add_picture(BytesIO(image_png_bytes), width=Cm(16))

    # 输出
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()
